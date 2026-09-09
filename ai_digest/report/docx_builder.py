"""Word 报告生成：双合集、逐条原文归档和可点击原文链接。"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

GROUPS = {
    "media": {
        "folder": "新闻媒体信息",
        "title": "新闻媒体机构报道摘录合集",
        "originals": "新闻报道原文",
    },
    "institution": {
        "folder": "机构信息",
        "title": "机构发布信息摘录合集",
        "originals": "机构发布原文",
    },
}
_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def collection_group(item: dict) -> str:
    """媒体来源单列；政府、智库、国际组织等统一归入机构类。"""
    return "media" if str(item.get("source_type", "")).lower() == "media" else "institution"


def safe_filename_component(value: str, max_chars: int) -> str:
    value = _INVALID_FILENAME.sub("_", str(value or ""))
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return (value or "未命名")[:max_chars].rstrip(" ._")


def original_filename(index: int, item: dict) -> str:
    source = safe_filename_component(item.get("source_name", "未知来源"), 30)
    title = safe_filename_component(item.get("zh_title") or item.get("title", "未命名"), 70)
    return f"{index:03d}_{source}_{title}.docx"


def _clean_text(value: object) -> str:
    return _INVALID_XML.sub("", str(value or ""))


def _set_cn_font(run, cn_name: str = "仿宋", size_pt: float = 16,
                 bold: bool = False, color: str | None = "000000") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), cn_name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _setup_document(doc: Document) -> None:
    """沿用现有 A4 政务报送版式，并显式固定正文样式。"""
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Pt(72)
        section.bottom_margin = Pt(72)
        section.left_margin = Pt(90)
        section.right_margin = Pt(90)
        section.header_distance = Mm(12.5)
        section.footer_distance = Mm(12.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(16)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(28)


def _add_para(doc: Document, text: object, *, cn_name: str = "仿宋",
              size_pt: float = 16, bold: bool = False, align=None,
              indent: bool = True, line_pt: float = 28, space_before: float = 0,
              space_after: float = 6, keep_with_next: bool = False):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_with_next
    if indent:
        fmt.first_line_indent = Pt(size_pt * 2)
    run = paragraph.add_run(_clean_text(text))
    _set_cn_font(run, cn_name=cn_name, size_pt=size_pt, bold=bold)
    return paragraph


def add_hyperlink(paragraph, url: str, text: str | None = None) -> None:
    """添加 Word 原生外部超链接，点击后交由系统默认浏览器访问。"""
    url = str(url or "").strip()
    if not url:
        return
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "仿宋")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    properties.extend([fonts, color, underline, size])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = _clean_text(text or url)
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_link_line(doc: Document, url: str) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(22)
    fmt.space_after = Pt(6)
    prefix = paragraph.add_run("原文链接：")
    _set_cn_font(prefix, size_pt=12)
    add_hyperlink(paragraph, url, text=url)


def _add_title_block(doc: Document, title: str, date_text: str,
                     window_text: str | None = None) -> None:
    _add_para(
        doc, title, cn_name="方正小标宋简体", size_pt=22, bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, line_pt=36,
        space_after=12, keep_with_next=True,
    )
    _add_para(
        doc, date_text, size_pt=16, align=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False, line_pt=28, space_after=4, keep_with_next=True,
    )
    if window_text:
        _add_para(
            doc, f"统计周期：{window_text}", size_pt=12,
            align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
            line_pt=22, space_after=18,
        )


def build_report_docx(items: Iterable[dict], out_path: str | Path,
                      title: str | None = None, date_text: str | None = None,
                      window_text: str | None = None) -> Path:
    """生成某一来源类别的摘要合集。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    values = list(items)
    doc = Document()
    _setup_document(doc)
    _add_title_block(
        doc,
        title or "人工智能安全治理信息摘报",
        date_text or datetime.now().strftime("%Y年%m月%d日"),
        window_text,
    )
    if not values:
        _add_para(doc, "本统计周期无入选信息。", indent=False)
    for index, item in enumerate(values, 1):
        _add_para(
            doc, f"{index}. {item.get('zh_title') or item.get('title', '')}",
            cn_name="黑体", size_pt=16, indent=False, line_pt=28,
            space_before=8, space_after=2, keep_with_next=True,
        )
        metadata = f"{item.get('source_name', '')} | {item.get('published_at', '')}"
        if item.get("category") or item.get("importance"):
            metadata += f" | {item.get('category', '')}/{item.get('importance', '')}"
        _add_para(
            doc, metadata, size_pt=12, indent=False, line_pt=22,
            space_after=2, keep_with_next=True,
        )
        _add_link_line(doc, item.get("url", ""))
        _add_para(doc, item.get("summary", ""), size_pt=16, line_pt=28, space_after=12)
    doc.save(str(out_path))
    return out_path


def build_original_docx(item: dict, out_path: str | Path, *, index: int) -> Path:
    """把一篇抓取原文独立保存为 Word 文件。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _setup_document(doc)
    title = item.get("title") or item.get("zh_title") or "未命名"
    _add_para(
        doc, f"{index}. {title}", cn_name="黑体", size_pt=20, bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, line_pt=32,
        space_after=12, keep_with_next=True,
    )
    _add_para(
        doc, f"出处：{item.get('source_name', '')}", size_pt=12,
        indent=False, line_pt=22, space_after=2, keep_with_next=True,
    )
    _add_para(
        doc, f"发布时间：{item.get('published_at', '')}", size_pt=12,
        indent=False, line_pt=22, space_after=2, keep_with_next=True,
    )
    _add_link_line(doc, item.get("url", ""))
    body = _clean_text(item.get("text", "")).strip()
    if not body:
        _add_para(doc, "未抓取到原文正文。", indent=False, size_pt=12, line_pt=22)
    else:
        for line in (line.strip() for line in body.splitlines()):
            if not line:
                continue
            for start in range(0, len(line), 2000):
                _add_para(doc, line[start:start + 2000], size_pt=12, line_pt=22, space_after=6)
    doc.save(str(out_path))
    return out_path


def build_report_bundle(items: Iterable[dict], output_dir: str | Path, *,
                        date_text: str | None = None,
                        date_stamp: str | None = None,
                        window_text: str | None = None) -> dict[str, dict]:
    """按媒体/机构生成两个合集，并为每篇文章保存独立原文 Word。"""
    output_dir = Path(output_dir)
    stamp = date_stamp or datetime.now().strftime("%Y%m%d")
    grouped = {"media": [], "institution": []}
    for item in items:
        grouped[collection_group(item)].append(item)

    bundle: dict[str, dict] = {}
    for key, spec in GROUPS.items():
        group_dir = output_dir / spec["folder"]
        originals_dir = group_dir / spec["originals"]
        originals_dir.mkdir(parents=True, exist_ok=True)
        collection_path = group_dir / f"{spec['title']}_{stamp}.docx"
        build_report_docx(
            grouped[key], collection_path, title=spec["title"],
            date_text=date_text, window_text=window_text,
        )
        originals = []
        for index, item in enumerate(grouped[key], 1):
            path = originals_dir / original_filename(index, item)
            originals.append(build_original_docx(item, path, index=index))
        bundle[key] = {
            "collection": collection_path,
            "originals": originals,
            "count": len(grouped[key]),
        }
    return bundle
